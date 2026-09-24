"""Civil ODS selections must survive the Telegram-to-document workflow."""

import io
import unittest
from unittest.mock import patch

import openpyxl
from pypdf import PdfReader

import app
import department_assumption_questions as questions
import department_contract_templates as templates
import department_ods_patch as departments


class CivilExclusionsTest(unittest.TestCase):
    def sample(self):
        return {
            'department': 'CIV', 'name': 'Client Exemple',
            'addr': '123 Rue Exemple\nMontréal QC', 'price': 2500,
            'odsNum': 'ODS26-097-CIV-JKV-Drainage',
            'service_lines': ['Conception du drainage;'],
        }

    def test_four_independent_exclusions_and_pdf_excel(self):
        data = self.sample()
        keys = ('civ_survey', 'civ_other_studies', 'civ_landscape', 'civ_supervision')
        for key in keys:
            data[f'dept_assumption_{key}_confirmed'] = True
            data[f'dept_assumption_{key}'] = True
        data['dept_assumption_civ_fees_confirmed'] = True
        data['dept_assumption_civ_fees'] = False
        data['dept_assumption_civ_docs'] = True
        selected = questions.selected_exclusions(data)
        self.assertEqual(len(selected), 4)
        self.assertFalse(any('permis' in item for item in selected))

        pdf = app.generate_pdf(data)
        pdf_text = '\n'.join(page.extract_text() for page in PdfReader(pdf).pages)
        for item in selected:
            self.assertIn(item.rstrip(';'), pdf_text)
        self.assertNotIn('Les études géotechniques, environnementales', pdf_text)

        excel = app.generate_excel(data)
        ws = openpyxl.load_workbook(io.BytesIO(excel.read()))['ODS']
        for item in selected:
            self.assertIn(item, ws['B55'].value)
        self.assertIn('Hypothèses', ws['B54'].value)

    def test_declined_exclusions_never_appear(self):
        data = self.sample()
        data['dept_assumption_civ_survey_confirmed'] = True
        data['dept_assumption_civ_survey'] = False
        self.assertEqual(questions.selected_exclusions(data), [])
        self.assertNotIn('Exclusions :', '\n'.join(
            page.extract_text() for page in PdfReader(app.generate_pdf(data)).pages
        ))
        self.assertEqual(templates.department_text(
            'Plans, relevés, certificat de localisation et informations existantes fournis avant le début du mandat, selon leur disponibilité;',
            'CIV', departments.DEPARTMENTS['CIV'],
        ).count('études géotechniques'), 0)

    def test_buttons_save_independent_yes_and_no_decisions(self):
        uid = '987654321'
        data = self.sample()
        data['special_note_confirmed'] = True
        app.user_data[uid] = data
        try:
            with patch.object(app, 'save_user_data'), patch.object(app, 'tg') as send, \
                 patch.object(app.req, 'post'):
                questions.ask_next_missing(100, uid)
                self.assertIn('civ_docs:yes', str(send.call_args))
                for key, answer in [('civ_docs', 'yes'), ('civ_survey', 'no')]:
                    questions.handle_update({'callback_query': {
                        'data': f'deptassume:{key}:{answer}', 'id': 'callback',
                        'from': {'id': int(uid)}, 'message': {'chat': {'id': 100}},
                    }})
                self.assertTrue(data['dept_assumption_civ_docs'])
                self.assertFalse(data['dept_assumption_civ_survey'])
                self.assertNotIn('Relevé d’arpentage;', questions.selected_exclusions(data))
        finally:
            app.user_data.pop(uid, None)

    def test_production_entrypoint_routes_civil_note_to_first_question(self):
        # Railway must run runtime:app; numbering_runtime:app bypasses this guard.
        from pathlib import Path
        import json
        import runtime

        railway = json.loads(Path('railway.json').read_text())
        self.assertIn('gunicorn runtime:app', railway['deploy']['startCommand'])
        uid = '987654322'
        data = self.sample()
        data.update({'special_note_confirmed': False, 'waiting_field': None})
        app.user_data[uid] = data
        try:
            with patch.object(app, 'save_user_data'), patch.object(app, 'tg') as send, \
                 patch.object(app.req, 'post'):
                runtime.department_dispatch({'callback_query': {
                    'data': 'note_none', 'id': 'callback',
                    'from': {'id': int(uid)}, 'message': {'chat': {'id': 100}},
                }})
                self.assertTrue(data['special_note_confirmed'])
                self.assertIn('deptassume:civ_docs:yes', str(send.call_args))
        finally:
            app.user_data.pop(uid, None)


if __name__ == '__main__':
    unittest.main()

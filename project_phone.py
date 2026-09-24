"""Invoice phone lookup by Project Code, with lossless XLSX cell updates.

Only the Phone cell is patched in the ZIP. No openpyxl save/round-trip, no
Description matching, no archive-phone fallback, and no workbook schema changes.
The Phone column is provisioned once in Excel before deploying this module.
"""
import io
import posixpath
import re
import urllib.parse
import zipfile
from xml.etree import ElementTree as ET
from xml.sax.saxutils import escape

import openpyxl


class PhoneStoreError(RuntimeError):
    pass


def project_code(value):
    """Extract the stable numbered code, ignoring the descriptive suffix."""
    match = re.match(r"^P(\d{2})-(\d{3})(?=$|[-_\s])", str(value or '').strip(), re.I)
    if not match:
        raise PhoneStoreError('Project Code invalide (exemple : P26-031).')
    return f'P{match[1]}-{match[2]}'


def _phone_text(value):
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def locate_phone(content, code):
    code = project_code(code)
    sheet = 'data 20' + code[1:3]
    workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=False)
    try:
        if sheet not in workbook.sheetnames:
            raise PhoneStoreError(f'Onglet {sheet} introuvable dans List.xlsx.')
        ws = workbook[sheet]
        rows = ws.iter_rows()
        headers = list(next(rows))
        columns = {}
        for label in ('Project Code', 'Phone'):
            found = [i for i, cell in enumerate(headers) if str(cell.value or '').strip() == label]
            if len(found) != 1:
                raise PhoneStoreError(f'La colonne {label} doit exister une seule fois dans {sheet}.')
            columns[label] = found[0]
        matches = []
        for row_number, row in enumerate(rows, 2):
            raw_code = row[columns['Project Code']].value
            try:
                key = project_code(raw_code)
            except PhoneStoreError:
                continue
            if key == code:
                matches.append((row_number, row[columns['Phone']]))
        if len(matches) != 1:
            reason = 'absent' if not matches else 'présent plusieurs fois'
            raise PhoneStoreError(f'{code} est {reason} dans la colonne Project Code de {sheet}.')
        row_number, cell = matches[0]
        if cell.data_type in ('f', 'e'):
            raise PhoneStoreError('La cellule Phone contient une formule ou une erreur. Vérifiez List.xlsx.')
        coordinate = openpyxl.utils.get_column_letter(columns['Phone'] + 1) + str(row_number)
        return sheet, coordinate, _phone_text(cell.value)
    finally:
        workbook.close()


def patch_phone(content, code, phone):
    phone = _phone_text(phone)
    if not phone or len(phone) > 100 or re.search(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', phone):
        raise PhoneStoreError('Numéro de téléphone vide ou invalide.')
    sheet, coordinate, existing = locate_phone(content, code)
    if existing:
        if existing == phone:
            return content
        raise PhoneStoreError('Phone a déjà été renseigné. Rouvrez Facturation pour lire sa valeur actuelle.')
    ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    rel_attr = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id'
    with zipfile.ZipFile(io.BytesIO(content)) as source:
        sheets = ET.fromstring(source.read('xl/workbook.xml')).find('s:sheets', ns)
        rel_id = next(s.attrib[rel_attr] for s in sheets if s.attrib['name'] == sheet)
        rels = ET.fromstring(source.read('xl/_rels/workbook.xml.rels'))
        target = next(r.attrib['Target'] for r in rels if r.attrib['Id'] == rel_id)
        path = target.lstrip('/') if target.startswith('/') else posixpath.normpath('xl/' + target)
        xml = source.read(path).decode('utf-8')
        row_number = re.search(r'\d+$', coordinate)[0]
        row_pattern = rf'(<row\b[^>]*\br="{row_number}"[^>]*>)(.*?)(</row>)'
        row_match = re.search(row_pattern, xml, re.S)
        if not row_match:
            raise PhoneStoreError('Format de ligne Excel non pris en charge. Aucune modification.')
        body = row_match[2]
        cell_pattern = rf'<c\b(?=[^>]*\br="{coordinate}")[^>]*(?:/>|>.*?</c>)'
        old = re.search(cell_pattern, body, re.S)
        style = ''
        if old:
            style_match = re.search(r'\bs="([^"]+)"', old[0])
            if style_match:
                style = f' s="{style_match[1]}"'
        cell = f'<c r="{coordinate}"{style} t="inlineStr"><is><t>{escape(phone)}</t></is></c>'
        if old:
            body = body[:old.start()] + cell + body[old.end():]
        else:
            col = openpyxl.utils.column_index_from_string(re.match(r'[A-Z]+', coordinate)[0])
            later = next((m for m in re.finditer(r'<c\b[^>]*\br="([A-Z]+)\d+"', body)
                          if openpyxl.utils.column_index_from_string(m[1]) > col), None)
            at = later.start() if later else len(body)
            body = body[:at] + cell + body[at:]
        xml = xml[:row_match.start(2)] + body + xml[row_match.end(2):]
        ET.fromstring(xml)  # Never upload malformed XML.
        output = io.BytesIO()
        with zipfile.ZipFile(output, 'w') as dest:
            dest.comment = source.comment
            for entry in source.infolist():
                dest.writestr(entry, xml.encode('utf-8') if entry.filename == path else source.read(entry))
    updated = output.getvalue()
    if locate_phone(updated, code)[2] != phone:
        raise PhoneStoreError('Vérification de Phone échouée. Aucune modification enregistrée.')
    return updated


class PhoneStore:
    def __init__(self, legacy):
        self.legacy = legacy

    def _connection(self):
        token = self.legacy.graph_access_token()
        sender = urllib.parse.quote(self.legacy.microsoft_email_config()['EMAIL_SENDER'], safe='')
        path = urllib.parse.quote(self.legacy.ODS_LIST_PATH, safe='/')
        return f'https://graph.microsoft.com/v1.0/users/{sender}/drive/root:/{path}', {
            'Authorization': f'Bearer {token}'
        }

    def _get(self, url, headers):
        response = self.legacy.req.get(url, headers=headers, timeout=40)
        if response.status_code != 200:
            raise PhoneStoreError(f'Lecture de List.xlsx impossible (HTTP {response.status_code}).')
        return response

    def read(self, code):
        with self.legacy.ods_list_lock:
            url, headers = self._connection()
            content = self._get(url + ':/content', headers).content
            return locate_phone(content, code)[2]

    def save(self, code, phone):
        # Share the existing ODS List lock: an ODS status update in this worker
        # must not overwrite a concurrent phone update.
        with self.legacy.ods_list_lock:
            url, headers = self._connection()
            metadata = self._get(url, headers).json()
            etag = metadata.get('eTag')
            if not etag:
                raise PhoneStoreError('Version de List.xlsx inconnue. Réessayez.')
            content = self._get(url + ':/content', headers).content
            if self._get(url, headers).json().get('eTag') != etag:
                raise PhoneStoreError('List.xlsx vient de changer. Renvoyez le numéro pour réessayer.')
            updated = patch_phone(content, code, phone)
            if updated == content:
                return
            response = self.legacy.req.put(
                url + ':/content',
                headers={**headers, 'If-Match': etag,
                         'Content-Type': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'},
                data=updated, timeout=40,
            )
            if response.status_code not in (200, 201):
                raise PhoneStoreError(f'Enregistrement de Phone impossible (HTTP {response.status_code}). Réessayez.')
            verified = self._get(url + ':/content', headers).content
            if locate_phone(verified, code)[2] != _phone_text(phone):
                raise PhoneStoreError('Phone non confirmé dans OneDrive. Renvoyez le numéro pour réessayer.')

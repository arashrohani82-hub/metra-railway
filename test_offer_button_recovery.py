"""A historical button must still identify its own offer after session changes."""
import io
import os
import tempfile
from unittest.mock import patch

import openpyxl
import pytest

os.environ.setdefault('ANTHROPIC_API_KEY', 'test-key')
os.environ.setdefault('DATA_DIR', tempfile.mkdtemp(prefix='offer-buttons-test-'))
import app
import ods_recovery as recovery
import offer_followup_runtime as followup
from offer_identity import reference, matches


def test_complete_ids_and_neighbouring_offers():
    assert reference('ODS26-135-STR-EDI — Examen') == 'ODS26-135-STR-EDI'
    assert reference('ODS26-1099-VDB') == 'ODS26-1099-VDB'
    assert not matches('ODS26-097-CIV-ABC', 'ODS26-097-INS')
    assert not matches('ODS26-0970-INS', 'ODS26-097')
    assert recovery._normalize_query('097') == 'ODS26-097'
    assert recovery._normalize_query('ODS26-097-INS') == 'ODS26-097-INS'


def test_old_conversion_button_uses_its_message_not_current_session():
    uid = '987654325'
    update = {'callback_query': {'data': 'of_convert', 'id': 'old', 'from': {'id': int(uid)},
        'message': {'chat': {'id': int(uid)}, 'text': '📬 ODS26-097-INS\n\nClient : Vincent Desjardins'}}}
    with patch.object(app, 'ALLOWED_USERS', {int(uid)}), \
         patch.object(app, 'user_data', {uid: {'offer_followup_selected': {'reference': 'ODS26-135-STR-EDI'}}}), \
         patch.object(followup, '_ack'), patch.object(app, 'tg'), patch.object(app.executor, 'submit') as submit:
        followup.handle_update_offer_followup(update)
    submit.assert_called_once_with(followup._confirm_conversion_from_followup, int(uid), uid,
                                   {'reference': 'ODS26-097-INS'})


def test_old_index_recovers_from_its_original_keyboard():
    callback = {'message': {'reply_markup': {'inline_keyboard': [[
        {'text': '🟠 48j · ODS26-097-INS · Vincent', 'callback_data': 'of_pick:1'}]]}}}
    assert followup._choice_reference(callback, '1') == 'ODS26-097-INS'
    assert followup._choice_reference({}, '1') == ''


def test_reference_button_survives_lost_session():
    uid = '987654325'
    offer = {'reference': 'ODS26-097-INS', 'date': '2026-08-19'}
    update = {'callback_query': {'data': 'of_pick:ODS26-097-INS', 'from': {'id': int(uid)},
                               'message': {'chat': {'id': int(uid)}}}}
    with patch.object(app, 'ALLOWED_USERS', {int(uid)}), patch.object(app, 'user_data', {}), \
         patch.object(followup, '_ack'), patch.object(app, 'save_user_data'), \
         patch.object(followup, 'load_open_offers', return_value=[offer]), \
         patch.object(followup, 'show_offer') as show:
        followup.handle_update_offer_followup(update)
        assert app.user_data[uid]['offer_followup_selected'] == offer
    show.assert_called_once_with(int(uid), uid)


def test_history_never_selects_another_suffix():
    records = {'ODS26-097-CIV-ABC': {'data': {'odsNum': 'ODS26-097-CIV-ABC'}}}
    with patch.object(app, 'offers_history', {'u': records}):
        assert followup._history_reference('u', 'ODS26-097-INS') == ''
        with patch.object(recovery, '_find_archived_xlsx', return_value=None) as find:
            assert recovery._recover_from_xlsx('u', 'ODS26-097-INS') is None
        find.assert_called_once_with('ODS26-097-INS')


def test_archived_ins_offer_is_restored_with_full_identity():
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'ODS'
    ws['B7'] = 'M. Vincent Desjardins'; ws['B10'] = 'Courriel : vdesjardins@gsf-canada.com'
    ws['B12'] = 'ODS26-097-INS'; ws['B47'] = 'Inspection structurale'; ws['E47'] = 2500
    raw = io.BytesIO(); wb.save(raw)
    filename = 'ODS26-097-INS-Inspection-structurale.xlsx'
    with patch.object(app, 'offers_history', {}), \
         patch.object(recovery, '_find_archived_xlsx', return_value=('token', 'owner', 'archive', filename)), \
         patch.object(app, 'download_onedrive_path', return_value=raw.getvalue()), \
         patch.object(app, 'save_offers_history'):
        restored = recovery._recover_from_xlsx('u', 'ODS26-097-INS')
        data = app.offers_history['u'][restored]['data']
        assert restored == 'ODS26-097-INS'
        assert data['price'] == 2500
        assert data['email'] == 'vdesjardins@gsf-canada.com'
        assert data['name'] == 'Vincent Desjardins'


def test_old_pending_button_recovers_before_confirmation():
    with patch.object(recovery, '_recover_from_xlsx', return_value='ODS26-097-INS') as restore, \
         patch.object(recovery, '_original_confirmation') as confirm:
        recovery.show_conversion_with_recovery(1, 'u', 'ODS26-097-INS')
    restore.assert_called_once_with('u', 'ODS26-097-INS')
    confirm.assert_called_once_with(1, 'u', 'ODS26-097-INS')


def test_ambiguous_base_reference_cannot_choose_a_project():
    records = {ref: {'data': {'odsNum': ref}} for ref in ('ODS26-097-INS', 'ODS26-097-CIV-ABC')}
    with patch.object(app, 'offers_history', {'u': records}), pytest.raises(ValueError, match='Plusieurs offres'):
        recovery._recover_from_xlsx('u', 'ODS26-097')

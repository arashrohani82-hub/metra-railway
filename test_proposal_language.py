"""English ODS integration: real documents, isolated provider/Telegram I/O."""
import copy
import io
import json
import re
import unittest
from concurrent.futures import ThreadPoolExecutor
from types import SimpleNamespace
from unittest.mock import patch

import openpyxl
from pypdf import PdfReader
import runtime
import app
import department_ods_patch as dept
import proposal_language as loc
from technical_content import parse_custom_technical_content


def fixture_translation(**kwargs):
    # Explicit deterministic test translations. No network or customer email.
    values = json.loads(kwargs['messages'][0]['content'])
    result = []
    for text in values:
        for k, v in sorted(loc._CACHE.items(), key=lambda p: -len(p[0])):
            if k.strip():
                text = text.replace(k, v)
        replacements = {
            'M./Mme ': '', 'M. ': 'Mr. ', 'Mme ': 'Ms. ',
            'Le délai de livraison estimé est de ': 'The estimated delivery time is ',
            ' jours ouvrables suivant la visite finale sur site.': ' business days following the final site visit.',
            'Photo ou texte du client…': 'Client photo or text…',
            '👋 Metra Consultation — Nouvelle offre': '👋 Metra Consultation - New proposal',
            'Inspections du drainage;': 'Drainage inspections;',
            'Analyse du drainage;': 'Drainage analysis;',
            'Génération Excel': 'Excel generation',
        }
        for k, v in replacements.items():
            text = text.replace(k, v)
        result.append(text)
    return SimpleNamespace(content=[SimpleNamespace(text=json.dumps(result))])


class ProposalLanguageTest(unittest.TestCase):
    def setUp(self):
        self.uid = '876543210'
        self.data = dict(language='en', province='ON', department='CIV',
                         name='Sample Client', civility='M./Mme', phone='416-555-0100',
                         email='sample@example.com', addr='123 Example Street\nToronto ON M1M 1M1',
                         price=6500, odsNum='ODS26-117-CIV-ABC-Drainage', date='2026-10-11',
                         delai='12', taxes='En sus', special_note='Keep access open',
                         service_lines=['Drainage analysis;', 'Permit and construction drawings;'],
                         project_title='Site drainage design')
        for key in ('civ_docs', 'civ_survey', 'civ_other_studies', 'civ_landscape', 'civ_supervision'):
            self.data[f'dept_assumption_{key}'] = True
            self.data[f'dept_assumption_{key}_confirmed'] = True
        self.key = patch.object(app.client.messages, 'create', side_effect=fixture_translation)
        self.provider = self.key.start()
        self.save = patch.object(app, 'save_user_data').start()
        loc.PREFERENCES.clear()
        loc.LANGUAGE.set('fr')

    def tearDown(self):
        patch.stopall()
        app.user_data.pop(self.uid, None)
        dept.DEPT_BY_UID.pop(self.uid, None)
        loc.PREFERENCES.clear()
        loc.LANGUAGE.set('fr')

    def test_english_pdf_and_excel_preserve_full_scope_and_selected_exclusions(self):
        self.data['service_lines'] = [f'Engineering service {i};' for i in range(1, 10)]
        original = copy.deepcopy(self.data)
        pdf = app.generate_pdf(self.data)
        text = '\n'.join(p.extract_text() for p in PdfReader(pdf).pages)
        for value in ('1. Description of services', 'HOURLY RATES', '25 %', '19.56 %',
                      'Senior engineer:', '130', '110', '105', '100', '85', '30 days',
                      'Project province: Ontario', 'Land survey', 'Geotechnical study',
                      'Landscape plans', 'Construction site supervision'):
            self.assertIn(value, text)
        for i in range(1, 10):
            self.assertIn(f'Engineering service {i}', text)
        self.assertNotIn('Les frais de permis', text)
        self.assertNotIn('HONORAIRES', text)
        ws = openpyxl.load_workbook(app.generate_excel(self.data))['ODS']
        self.assertEqual(ws['F47'].value, '=E47*D47')
        self.assertEqual(ws['E47'].value, 6500)
        self.assertIn('12 business days', ws['B51'].value)
        self.assertIn('Keep access open', ws['B56'].value)
        self.assertIn('Extra', ws['B50'].value)
        self.assertEqual(ws['B7'].value, 'Sample Client')
        for value in ('Land survey', 'Geotechnical study', 'Landscape plans', 'Construction site supervision'):
            self.assertIn(value, ws['B55'].value)
        self.assertIn('Civil Engineering', ws['B1'].value)
        self.assertEqual(self.data, original)

    def test_three_departments_and_both_languages_are_isolated(self):
        def render(code, language):
            data = dict(self.data, department=code, language=language)
            return '\n'.join(p.extract_text() for p in PdfReader(app.generate_pdf(data)).pages)
        with ThreadPoolExecutor(max_workers=3) as pool:
            en, fr, geo = list(pool.map(lambda args: render(*args), [('STR','en'), ('STR','fr'), ('GEO','en')]))
        self.assertIn('Structural Engineering', en)
        self.assertIn('HONORAIRES', fr)
        self.assertNotIn('HONORAIRES', en)
        self.assertIn('Geotechnical Engineering', geo)
        self.assertIn('soil and rock conditions', geo)

    def test_outgoing_email_matches_english_preview_and_department(self):
        recipient, subject, body = app.build_email_preview(self.data)
        self.assertEqual(recipient, 'sample@example.com')
        self.assertIn('Service proposal ODS26-117-CIV-ABC', subject)
        self.assertIn(self.data['addr'], body)
        rendered = app.email_body_html(body)
        self.assertIn('President - Civil Engineer', rendered)
        self.assertNotIn('Cordialement', rendered)
        self.assertEqual(rendered.count('Kind regards,'), 1)
        self.assertIn('metraconsultation.ca', rendered)

    def test_source_draft_unchanged_and_no_provider_for_french(self):
        data = dict(self.data, language='fr')
        app.generate_pdf(data)
        app.generate_excel(data)
        app.build_email_preview(data)
        self.provider.assert_not_called()

    def callback(self, value):
        return {'callback_query': {'id': 'cb', 'data': value, 'from': {'id': int(self.uid)}, 'message': {'chat': {'id': int(self.uid)}}}}

    def test_language_province_selection_survives_department_reset_and_extraction(self):
        with patch.object(app, 'ALLOWED_USERS', {int(self.uid)}), patch.object(app, 'tg'), patch.object(app.req, 'post'):
            runtime.department_dispatch(self.callback('ods_lang:en'))
            runtime.department_dispatch(self.callback('ods_province:BC'))
            runtime.department_dispatch(self.callback('dept:GEO'))
            self.assertEqual(app.user_data[self.uid]['language'], 'en')
            self.assertEqual(app.user_data[self.uid]['province'], 'BC')
            app.user_data[self.uid] = {'department':'GEO', 'name':'Extracted Client'}
            with patch.object(loc, 'translate_many', side_effect=lambda xs, protected=(): xs), patch.object(app, 'ask_desc_options'), patch.object(app, 'get_next_project_num', return_value='117'):
                app.ask_next_missing(int(self.uid), self.uid)
            self.assertEqual(app.user_data[self.uid]['language'], 'en')
            self.assertEqual(app.user_data[self.uid]['province'], 'BC')
            self.assertEqual(app.user_data[self.uid]['department'], 'GEO')

    def test_unauthorized_language_callback_cannot_change_state(self):
        with patch.object(app, 'ALLOWED_USERS', set()), patch.object(app, 'tg'):
            runtime.department_dispatch(self.callback('ods_lang:en'))
        self.assertNotIn(self.uid, app.user_data)

    def test_changed_numbers_rejected_and_not_cached(self):
        phrase = 'Forage 37 mètres'
        response = SimpleNamespace(content=[SimpleNamespace(text='["Borehole 38 metres"]')])
        loc._CACHE.pop(phrase, None)
        with patch.object(app.client.messages, 'create', return_value=response):
            with self.assertRaisesRegex(ValueError, 'number'):
                loc.translate_many([phrase])
        self.assertNotIn(phrase, loc._CACHE)

    def test_incomplete_translation_rejected(self):
        response = SimpleNamespace(content=[SimpleNamespace(text='[]')])
        with patch.object(app.client.messages, 'create', return_value=response):
            with self.assertRaisesRegex(ValueError, 'Incomplete'):
                loc.translate_many(['Donnée unique à traduire'])

    def test_executor_propagates_language(self):
        token = loc.LANGUAGE.set('en')
        try:
            self.assertEqual(app.executor.submit(loc.LANGUAGE.get).result(timeout=3), 'en')
        finally:
            loc.LANGUAGE.reset(token)

    def test_english_scope_headers(self):
        mandate, lines = parse_custom_technical_content('Mandate: Assess the floor\nServices:\n• Site inspection\n• Written report')
        self.assertEqual(mandate, 'Assess the floor')
        self.assertEqual(len(lines), 2)

    def test_offer_history_keeps_language_and_province(self):
        record = app.history_data_copy(self.data)
        self.assertEqual(record['language'], 'en')
        self.assertEqual(record['province'], 'ON')

    def test_english_followup_uses_saved_offer_language_and_one_signature(self):
        import offer_followup_runtime as followups
        reference = 'ODS26-117-CIV-ABC'
        app.offers_history[self.uid] = {reference: {'data': dict(self.data)}}
        try:
            with patch.object(loc, 'translate_many', return_value=['Follow-up - ' + reference, 'Hello Sample Client,\n\nFollowing up on your proposal.']):
                subject, body = followups.build_followup_email({'reference': reference, 'client': 'Sample Client'}, 7)
            self.assertIn('Follow-up', subject)
            rendered = followups._followup_html(body)
            self.assertEqual(rendered.count('Kind regards,'), 1)
            self.assertNotIn('Bien cordialement', rendered)
            self.assertNotIn('Président', rendered)
            self.assertIn('a.rohani@metraconsultation.ca', rendered)
        finally:
            app.offers_history.pop(self.uid, None)


if __name__ == '__main__':
    unittest.main()

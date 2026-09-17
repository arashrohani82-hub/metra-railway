"""Recover archived ODS offers when local Railway history was lost after deploys."""
import io
import re
from datetime import datetime

import openpyxl
import app as ods

BASE_ARCHIVE_ROOTS = (
    "Metra Structure Inc/Offre de service",
    "Metra Consultation Inc/Offre de service",
    "Metra Consultation/Offre de service",
)
DEPARTMENT_FOLDERS = ("Structure", "Civil", "Geotechnic")
_original_show_pending_offers = ods.show_pending_offers


def _normalize_query(query):
    text = str(query or "").strip().upper()
    digits = re.search(r"(?:ODS(\d{2})-)?(\d{1,4})", text)
    if not digits:
        return text
    year = digits.group(1) or datetime.now().strftime('%y')
    return f"ODS{year}-{int(digits.group(2)):03d}"


def _cell_text(value, prefix=""):
    text = str(value or "").strip()
    if prefix and text.lower().startswith(prefix.lower()):
        text = text[len(prefix):].strip()
    return text


def _safe_price(value):
    if value is None:
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace("$", "").replace(" ", "")
    if "," in text and "." not in text:
        text = text.replace(",", ".")
    else:
        text = text.replace(",", "")
    try:
        return float(text)
    except Exception:
        return 0.0


def _archive_roots():
    """Current three-folder layout first, then legacy locations."""
    year = datetime.now().strftime('%Y')
    roots = []
    for base in BASE_ARCHIVE_ROOTS:
        roots.extend(f"{base}/{folder}" for folder in DEPARTMENT_FOLDERS)
        roots.append(base)
        roots.extend([
            f"{base}/{year}/Offres Structure",
            f"{base}/{year}/Offres Civil",
            f"{base}/{year}/Offres Géotechnique",
        ])
    # Preserve order while removing duplicates.
    return tuple(dict.fromkeys(roots))


def _find_archived_xlsx(ref):
    token = ods.graph_access_token()
    sender = ods.microsoft_email_config()["EMAIL_SENDER"]
    errors = []
    for root in _archive_roots():
        try:
            items = ods.list_onedrive_children(token, sender, root)
        except Exception as exc:
            errors.append(f"{root}: {exc}")
            continue
        candidates = []
        for item in items:
            name = str(item.get("name") or "")
            if name.lower().endswith(".xlsx") and ref.casefold() in name.casefold():
                candidates.append(name)
        if candidates:
            filename = sorted(
                candidates,
                key=lambda n: (0 if n.upper().startswith(ref.upper()) else 1, len(n)),
            )[0]
            return token, sender, root, filename
    if errors:
        ods.logger.warning("ODS recovery roots checked with errors: %s", " | ".join(errors))
    return None


def _recover_from_xlsx(uid, query):
    ref = _normalize_query(query)
    if not re.fullmatch(r"ODS\d{2}-\d{3,4}", ref, re.I):
        return None

    records = ods.offers_history.get(str(uid), {})
    for existing_ref in records:
        if ref.casefold() in str(existing_ref).casefold():
            return existing_ref

    found = _find_archived_xlsx(ref)
    if not found:
        ods.logger.warning("ODS RECOVERY NOT FOUND IN ONEDRIVE: %s", ref)
        return None
    token, sender, root, filename = found
    raw = ods.download_onedrive_path(token, sender, f"{root}/{filename}")

    wb_values = openpyxl.load_workbook(io.BytesIO(raw), data_only=True)
    ws = wb_values["ODS"] if "ODS" in wb_values.sheetnames else wb_values[wb_values.sheetnames[0]]
    wb_formula = openpyxl.load_workbook(io.BytesIO(raw), data_only=False)
    wsf = wb_formula["ODS"] if "ODS" in wb_formula.sheetnames else wb_formula[wb_formula.sheetnames[0]]

    def value_at(cell):
        value = ws[cell].value
        if value in (None, ""):
            fallback = wsf[cell].value
            if not (isinstance(fallback, str) and fallback.startswith("=")):
                value = fallback
        return value

    identity = _cell_text(value_at("B7"))
    civility = "M./Mme"
    name = identity
    for prefix, value in (("Mme ", "Mme"), ("M. ", "M."), ("M ", "M.")):
        if identity.startswith(prefix):
            civility = value
            name = identity[len(prefix):].strip()
            break

    # New format: ODS26-123-CIV-ABC. Old formats remain supported.
    full_pattern = r"ODS\d{2}-\d{3,4}-(?:STR|CIV|GEO)-[A-Z]{3}"
    old_pattern = r"ODS\d{2}-\d{3,4}(?:-[A-Z]{3})?"
    filename_match = re.search(full_pattern, filename, re.I) or re.search(old_pattern, filename, re.I)
    ods_cell = _cell_text(value_at("B12"))
    cell_match = re.search(full_pattern, ods_cell, re.I) or re.search(old_pattern, ods_cell, re.I)
    ods_match = filename_match or cell_match
    ods_num = ods_match.group(0).upper() if ods_match else ref

    department = "STR"
    file_code = "PRJ"
    current = re.search(r"ODS\d{2}-\d{3,4}-(STR|CIV|GEO)-([A-Z]{3})", ods_num, re.I)
    if current:
        department = current.group(1).upper()
        file_code = current.group(2).upper()
    else:
        # Folder is a reliable fallback for the department in the new layout.
        root_lower = root.casefold()
        if root_lower.endswith('/civil'):
            department = 'CIV'
        elif root_lower.endswith('/geotechnic') or root_lower.endswith('/offres géotechnique'):
            department = 'GEO'
        legacy_code = re.search(r"ODS\d{2}-\d{3,4}-([A-Z]{3})", ods_num, re.I)
        if legacy_code:
            file_code = legacy_code.group(1).upper()

    desc = _cell_text(value_at("B47"))
    price = _safe_price(value_at("E47"))
    data = {
        "name": ods.normalize_client_name(name),
        "civility": civility,
        "phone": _cell_text(value_at("B9"), "Cell. :"),
        "email": _cell_text(value_at("B10"), "Courriel :"),
        "addr": _cell_text(value_at("B8"), "Adresse :"),
        "desc": desc,
        "service": desc,
        "project_title": desc.split("\n", 1)[0][:120] if desc else "Projet",
        "price": price,
        "odsNum": ods_num,
        "department": department,
        "file_code": file_code,
        "date": datetime.now().strftime("%Y-%m-%d"),
        "email_sent_at": datetime.now().isoformat(timespec="seconds"),
        "recovered_from_onedrive": True,
        "recovered_archive_file": filename,
        "recovered_archive_root": root,
    }
    recovered_ref = ods.record_sent_offer(str(uid), data)
    ods.logger.warning(
        "ODS RECOVERED FROM ONEDRIVE: %s file=%s root=%s department=%s price=%s",
        recovered_ref, filename, root, department, price,
    )
    return recovered_ref


def show_pending_offers_with_recovery(chat_id, uid, query=""):
    if query and not ods.pending_offer_records(uid, query):
        try:
            recovered = _recover_from_xlsx(uid, query)
            if recovered:
                query = _normalize_query(query)
        except Exception as exc:
            ods.logger.exception("ODS OneDrive recovery failed: %s", exc)
    return _original_show_pending_offers(chat_id, uid, query)


ods.show_pending_offers = show_pending_offers_with_recovery
ods.logger.warning("ODS ONEDRIVE RECOVERY ACTIVE v3: department folders + full STR/CIV/GEO refs")
